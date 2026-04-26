from __future__ import annotations

from io import BytesIO
import unittest

from skillproof.ai_assist import assessment_payload, build_reviewer_prompt
from skillproof.assessment import build_assessment, score_assessment
from skillproof.file_readers import read_file_bytes
from skillproof.report import build_markdown_report, learning_plan_rows
from skillproof.sample_data import SAMPLE_ANSWERS, SAMPLE_JD, SAMPLE_RESUME


class SkillProofTests(unittest.TestCase):
    def test_extracts_core_skills_from_sample(self) -> None:
        assessment = build_assessment(SAMPLE_JD, SAMPLE_RESUME)
        names = {skill.name for skill in assessment.skills}
        self.assertIn("Python", names)
        self.assertIn("React", names)
        self.assertIn("Testing", names)

    def test_scores_and_generates_report(self) -> None:
        assessment = build_assessment(SAMPLE_JD, SAMPLE_RESUME)
        scored = score_assessment(assessment, SAMPLE_ANSWERS)
        self.assertGreater(scored.overall_score, 50)
        report = build_markdown_report(scored)
        self.assertIn("Overall readiness", report)
        self.assertIn("Personalized Learning Plan", report)

    def test_learning_plan_rows_include_gaps(self) -> None:
        assessment = build_assessment(SAMPLE_JD, SAMPLE_RESUME)
        scored = score_assessment(assessment, SAMPLE_ANSWERS)
        plans = learning_plan_rows(scored)
        self.assertTrue(plans)
        self.assertIn("Timeline", plans[0])

    def test_reads_csv_as_structured_evidence(self) -> None:
        content = b"name,skills,impact\nSehaj,Python and React,Automated 40 hours of review\n"
        text = read_file_bytes("candidate.csv", content)
        self.assertIn("skills: Python and React", text)
        self.assertIn("impact: Automated 40 hours of review", text)

    def test_reads_docx_and_xlsx_files(self) -> None:
        from docx import Document
        from openpyxl import Workbook

        docx_buffer = BytesIO()
        document = Document()
        document.add_paragraph("Built Python APIs and React dashboards.")
        document.save(docx_buffer)

        xlsx_buffer = BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ATS export"
        sheet.append(["Skill", "Evidence"])
        sheet.append(["Testing", "Wrote regression tests for scoring engine"])
        workbook.save(xlsx_buffer)

        self.assertIn("Python APIs", read_file_bytes("resume.docx", docx_buffer.getvalue()))
        self.assertIn("Skill: Testing", read_file_bytes("candidate.xlsx", xlsx_buffer.getvalue()))

    def test_ai_reviewer_prompt_uses_structured_payload(self) -> None:
        assessment = build_assessment(SAMPLE_JD, SAMPLE_RESUME)
        scored = score_assessment(assessment, SAMPLE_ANSWERS)
        payload = assessment_payload(scored, "Project-based", 6)
        prompt = build_reviewer_prompt(scored, "Project-based", 6)

        self.assertIn("overall_score", payload)
        self.assertIn("Business Impact Story", prompt)
        self.assertNotIn(SAMPLE_RESUME[:80], prompt)
        self.assertNotIn(SAMPLE_JD[:80], prompt)


if __name__ == "__main__":
    unittest.main()
